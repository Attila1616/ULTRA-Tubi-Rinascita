import json
import os
import re
import shutil
from collections import defaultdict
from datetime import datetime

import openpyxl


SCHEMA_VERSION = 1
CODE_CATALOG_SCHEMA_VERSION = 2
DATABASE_DIR_NAME = "Database Tubi"
CODE_SOURCE_DIR_NAME = "Codici Tubi"
MANIFEST_NAME = "manifest.json"
MOVEMENTS_NAME = "movimenti.jsonl"
CODE_CATALOG_NAME = "codici_tubi.json"
ADHOC_STATE_NAME = "adhoc_state.json"
ADHOC_LOG_NAME = "adhoc_reset_log.jsonl"

HEADERS = ["Misura", "Spessore", "Materiale", "Finitura", "Quantita"]

_CONFIGURED_DATABASE_DIR = None
_CONFIGURED_CODE_CATALOG_PATH = None


def configure_storage(database_path=None, code_catalog_file=None):
    """Configure shared storage paths while retaining project-local defaults."""
    global _CONFIGURED_DATABASE_DIR, _CONFIGURED_CODE_CATALOG_PATH
    database_path = str(database_path or "").strip()
    code_catalog_file = str(code_catalog_file or "").strip()
    _CONFIGURED_DATABASE_DIR = os.path.abspath(database_path) if database_path else None
    _CONFIGURED_CODE_CATALOG_PATH = os.path.abspath(code_catalog_file) if code_catalog_file else None
    return {
        "database_dir": _CONFIGURED_DATABASE_DIR,
        "code_catalog_path": _CONFIGURED_CODE_CATALOG_PATH,
    }


def database_dir(project_root):
    return _CONFIGURED_DATABASE_DIR or os.path.join(project_root, DATABASE_DIR_NAME)


def manifest_path(project_root):
    return os.path.join(database_dir(project_root), MANIFEST_NAME)


def movements_path(project_root):
    return os.path.join(database_dir(project_root), MOVEMENTS_NAME)


def code_source_dir(project_root):
    return os.path.join(project_root, CODE_SOURCE_DIR_NAME)


def code_catalog_path(project_root):
    return _CONFIGURED_CODE_CATALOG_PATH or os.path.join(project_root, CODE_CATALOG_NAME)


def legacy_code_catalog_path(project_root):
    return os.path.join(database_dir(project_root), CODE_CATALOG_NAME)


def adhoc_state_path(project_root):
    return os.path.join(database_dir(project_root), ADHOC_STATE_NAME)


def adhoc_log_path(project_root):
    return os.path.join(database_dir(project_root), ADHOC_LOG_NAME)


def normalize_decimal_text(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        if float(value).is_integer():
            return str(int(value))
        return f"{float(value):.6f}".rstrip("0").rstrip(".")
    text = str(value).strip().replace(",", ".")
    if not text:
        return ""
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
        return f"{number:.6f}".rstrip("0").rstrip(".")
    except ValueError:
        return text


def decimal_number(value):
    text = normalize_decimal_text(value)
    if not text:
        return None
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


def normalize_misura(value):
    text = normalize_decimal_text(value).replace(" ", "")
    text = text.replace("X", "x")
    parts = [normalize_decimal_text(part) for part in text.split("x") if part != ""]
    return "x".join(parts) if len(parts) > 1 else text


def normalize_materiale(value):
    return str(value or "").strip().upper()


def normalize_finitura(value):
    return str(value or "").strip().upper() or "2B"


def parse_quantity(value):
    if value is None or value == "":
        return 0
    try:
        return int(float(str(value).strip().replace(",", ".")))
    except ValueError:
        return 0


def classify_tube(misura):
    misura = normalize_misura(misura)
    if "x" not in misura:
        return "tondo", [decimal_number(misura)]

    dimensions = [decimal_number(part) for part in misura.split("x")]
    if len(dimensions) >= 2 and dimensions[0] == dimensions[1]:
        return "quadrato", dimensions[:2]
    return "rettangolare", dimensions[:2]


def tube_key(tube):
    return (
        normalize_misura(tube.get("misura")),
        normalize_decimal_text(tube.get("spessore")),
        normalize_materiale(tube.get("materiale")),
        normalize_finitura(tube.get("finitura")),
    )


def make_tube(misura, spessore, materiale, finitura, quantita):
    misura_norm = normalize_misura(misura)
    forma, dimensioni = classify_tube(misura_norm)
    return {
        "misura": misura_norm,
        "forma": forma,
        "dimensioni": dimensioni,
        "spessore": decimal_number(spessore),
        "materiale": normalize_materiale(materiale),
        "finitura": normalize_finitura(finitura),
        "quantita": parse_quantity(quantita),
    }


def identity_key_from_values(misura, spessore, materiale, finitura):
    return "|".join([
        normalize_misura(misura),
        normalize_decimal_text(spessore),
        normalize_materiale(materiale) or "304",
        normalize_finitura(finitura) or "2B",
    ])


def identity_key_from_tube(tube):
    return identity_key_from_values(
        tube.get("misura"),
        tube.get("spessore"),
        tube.get("materiale"),
        tube.get("finitura"),
    )


def identity_dict_from_key(key):
    parts = str(key or "").split("|")
    while len(parts) < 4:
        parts.append("")
    misura, spessore, materiale, finitura = parts[:4]
    forma, dimensioni = classify_tube(misura)
    return {
        "key": key,
        "misura": misura,
        "spessore": spessore,
        "materiale": materiale or "304",
        "finitura": finitura or "2B",
        "forma": forma,
        "dimensioni": dimensioni,
    }


def display_tube_measure(misura, spessore, materiale=None, finitura=None, include_defaults=False):
    misura = normalize_misura(misura)
    spessore = normalize_decimal_text(spessore)
    prefix = "Ø" if misura and "x" not in misura else ""
    label = f"{prefix}{misura}x{spessore}" if spessore else f"{prefix}{misura}"
    mat = normalize_materiale(materiale) or "304"
    fin = normalize_finitura(finitura) or "2B"
    if include_defaults or mat != "304" or fin != "2B":
        label = f"{label} {mat} {fin}"
    return label


def _extract_codes_from_text(text):
    return [code.upper() for code in re.findall(r"\bTA\d{4}A\d+\b", str(text or ""), flags=re.IGNORECASE)]


def _clean_thickness_text(value):
    text = normalize_decimal_text(value)
    text = re.sub(r"(?i)\b(304|316|2b|ba|lucido|fe)\b", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return normalize_decimal_text(text)


def _sort_number(value):
    try:
        return float(normalize_decimal_text(value))
    except (TypeError, ValueError):
        return 0.0


def _material_from_text(text):
    return "316" if re.search(r"\b316\b", str(text or ""), flags=re.IGNORECASE) else "304"


def _finish_from_text(text):
    return "BA" if re.search(r"\b(BA|LUCIDO)\b", str(text or ""), flags=re.IGNORECASE) else "2B"


def _preferred_code_from_row(row_values, code_column_index):
    if len(row_values) > code_column_index:
        column_codes = _extract_codes_from_text(row_values[code_column_index])
        if column_codes:
            return column_codes[0]
    joined_codes = _extract_codes_from_text(" ".join(str(value or "") for value in row_values))
    return joined_codes[0] if joined_codes else ""


def _merge_code_entry(entries, misura, spessore, materiale, finitura, code, source_name, row_number):
    code = str(code or "").strip().upper()
    if not code:
        return
    key = identity_key_from_values(misura, spessore, materiale, finitura)
    if key not in entries:
        identity = identity_dict_from_key(key)
        entries[key] = {
            "key": key,
            "misura": identity["misura"],
            "spessore": identity["spessore"],
            "materiale": identity["materiale"],
            "finitura": identity["finitura"],
            "forma": identity["forma"],
            "code": code,
            "sources": [],
        }
    source_ref = {"file": source_name, "row": row_number}
    if source_ref not in entries[key]["sources"]:
        entries[key]["sources"].append(source_ref)


def _parse_round_code_row(row_values, source_name, row_number):
    joined = " ".join(str(v or "") for v in row_values)
    code = _preferred_code_from_row(row_values, 3)
    if not code:
        return None

    misura = row_values[1] if len(row_values) > 1 else None
    spessore = row_values[2] if len(row_values) > 2 else None
    if not normalize_misura(misura) or not _clean_thickness_text(spessore):
        return None

    return {
        "misura": normalize_misura(misura),
        "spessore": _clean_thickness_text(spessore),
        "materiale": _material_from_text(joined),
        "finitura": _finish_from_text(joined),
        "code": code,
        "source": source_name,
        "row": row_number,
    }


def _parse_tubolare_code_row(row_values, source_name, row_number):
    joined = " ".join(str(v or "") for v in row_values)
    code = _preferred_code_from_row(row_values, 2)
    if not code:
        return None

    size_match = re.search(r"(\d+(?:[.,]\d+)?x\d+(?:[.,]\d+)?x\d+(?:[.,]\d+)?)", joined, flags=re.IGNORECASE)
    if not size_match:
        return None

    parts = [normalize_decimal_text(part) for part in size_match.group(1).lower().replace(",", ".").split("x")]
    if len(parts) < 3:
        return None

    return {
        "misura": normalize_misura("x".join(parts[:2])),
        "spessore": _clean_thickness_text(parts[2]),
        "materiale": _material_from_text(joined),
        "finitura": _finish_from_text(joined),
        "code": code,
        "source": source_name,
        "row": row_number,
    }


def is_location_sheet(sheet_name):
    name = str(sheet_name or "").strip().lower()
    if not name:
        return False
    if name.startswith("accoda"):
        return False
    if "conteggio" in name or "riepilogo" in name or "totale" in name:
        return False
    return name == "terra" or name.startswith("scaffalatura")


def location_type(sheet_name):
    return "terra" if str(sheet_name).strip().lower() == "terra" else "scaffalatura"


def safe_location_filename(location_name):
    safe = re.sub(r'[<>:"/\\|?*]+', " ", str(location_name).strip())
    safe = re.sub(r"\s+", " ", safe).strip()
    return f"{safe}.json"


def header_matches(values):
    normalized = [str(v or "").strip().lower().replace("à", "a").replace(" ", "") for v in values]
    return normalized == ["misura", "spessore", "materiale", "finitura", "quantita"]


def merged_title_for_cell(ws, row, col):
    for merged_range in ws.merged_cells.ranges:
        if ws.cell(row=row, column=col).coordinate in merged_range:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
    return ws.cell(row=row, column=col).value


def find_tables(ws):
    tables = []
    for row in range(1, (ws.max_row or 0) + 1):
        for col in range(1, (ws.max_column or 0) - 3):
            values = [ws.cell(row=row, column=col + offset).value for offset in range(5)]
            if not header_matches(values):
                continue
            title = merged_title_for_cell(ws, row - 1, col) if row > 1 else None
            title = str(title or "").strip()
            if not title:
                title = "Ripiano 1"
            tables.append({"name": title, "header_row": row, "start_col": col})
    return tables


def merge_tube_into_list(tubes, tube):
    if tube["quantita"] <= 0:
        return

    key = tube_key(tube)
    for existing in tubes:
        if tube_key(existing) == key:
            existing["quantita"] += tube["quantita"]
            return

    tubes.append(tube)


def sort_tubes(tubes):
    return sorted(
        tubes,
        key=lambda tube: (
            tube.get("forma", ""),
            str(tube.get("misura", "")),
            float(tube.get("spessore") or 0),
            str(tube.get("materiale", "")),
            str(tube.get("finitura", "")),
        ),
    )


def parse_location_sheet(ws):
    ripiani = []

    for table in find_tables(ws):
        tubes = []
        row = table["header_row"] + 1
        col = table["start_col"]

        while row <= (ws.max_row or 0):
            misura = ws.cell(row=row, column=col).value
            spessore = ws.cell(row=row, column=col + 1).value

            if normalize_misura(misura) == "" and normalize_decimal_text(spessore) == "":
                break

            tube = make_tube(
                misura,
                spessore,
                ws.cell(row=row, column=col + 2).value,
                ws.cell(row=row, column=col + 3).value,
                ws.cell(row=row, column=col + 4).value,
            )
            merge_tube_into_list(tubes, tube)
            row += 1

        ripiani.append({
            "name": table["name"],
            "tubi": sort_tubes(tubes),
        })

    return ripiani


def write_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write("\n")


def read_json(path, default):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return default


def empty_code_catalog():
    return {
        "schema_version": CODE_CATALOG_SCHEMA_VERSION,
        "updated_at": None,
        "sources": [],
        "entries": [],
    }


def _entry_code(entry):
    code = str((entry or {}).get("code") or "").strip().upper()
    if code:
        return code
    legacy_codes = [str(value or "").strip().upper() for value in (entry or {}).get("codes", []) if str(value or "").strip()]
    return legacy_codes[0] if legacy_codes else ""


def _normalize_code_catalog(catalog):
    normalized = {
        **(catalog or {}),
        "schema_version": CODE_CATALOG_SCHEMA_VERSION,
        "sources": list((catalog or {}).get("sources", [])),
        "entries": [],
    }
    for entry in (catalog or {}).get("entries", []):
        code = _entry_code(entry)
        if not code:
            continue
        clean_entry = {key: value for key, value in entry.items() if key != "codes"}
        clean_entry["code"] = code
        normalized["entries"].append(clean_entry)
    return normalized


def load_code_catalog(project_root):
    path = code_catalog_path(project_root)
    if not os.path.exists(path) and os.path.exists(legacy_code_catalog_path(project_root)):
        path = legacy_code_catalog_path(project_root)
    return _normalize_code_catalog(read_json(path, empty_code_catalog()))


def save_code_catalog(project_root, catalog):
    catalog = _normalize_code_catalog(catalog)
    catalog["updated_at"] = datetime.now().isoformat(timespec="seconds")
    write_json(code_catalog_path(project_root), catalog)


def import_codes_from_excel(project_root, source_dir=None):
    source_dir = source_dir or code_source_dir(project_root)
    if not os.path.isdir(source_dir):
        return {"status": "error", "message": f"Cartella codici non trovata: {source_dir}"}

    entries = {}
    sources = []
    rows_read = 0
    rows_imported = 0

    for filename in sorted(os.listdir(source_dir)):
        if filename.startswith("~$") or not filename.lower().endswith((".xlsx", ".xlsm")):
            continue
        path = os.path.join(source_dir, filename)
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sources.append({"file": filename, "mtime": datetime.fromtimestamp(os.path.getmtime(path)).isoformat(timespec="seconds")})
        is_round_source = "tubi" in filename.lower() and "tubolari" not in filename.lower()

        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                row_values = list(row)
                if not any(value is not None and str(value).strip() for value in row_values):
                    continue
                rows_read += 1
                parsed = _parse_round_code_row(row_values, filename, row_number) if is_round_source else _parse_tubolare_code_row(row_values, filename, row_number)
                if not parsed:
                    continue
                _merge_code_entry(
                    entries,
                    parsed["misura"],
                    parsed["spessore"],
                    parsed["materiale"],
                    parsed["finitura"],
                    parsed["code"],
                    parsed["source"],
                    parsed["row"],
                )
                rows_imported += 1

    old_catalog = load_code_catalog(project_root)
    for old_entry in old_catalog.get("entries", []):
        if not old_entry.get("manual"):
            continue
        key = identity_key_from_values(old_entry.get("misura"), old_entry.get("spessore"), old_entry.get("materiale"), old_entry.get("finitura"))
        code = _entry_code(old_entry)
        if not code:
            continue
        identity = identity_dict_from_key(key)
        entries[key] = {
            **old_entry,
            "key": key,
            "misura": identity["misura"],
            "spessore": identity["spessore"],
            "materiale": identity["materiale"],
            "finitura": identity["finitura"],
            "forma": identity["forma"],
            "code": code,
            "manual": True,
        }

    catalog = {
        "schema_version": CODE_CATALOG_SCHEMA_VERSION,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "sources": sources,
        "entries": sorted(entries.values(), key=lambda item: (
            item.get("forma", ""),
            item.get("misura", ""),
            _sort_number(item.get("spessore")),
            item.get("materiale", ""),
            item.get("finitura", ""),
        )),
    }
    save_code_catalog(project_root, catalog)
    return {
        "status": "success",
        "catalog_path": code_catalog_path(project_root),
        "sources_count": len(sources),
        "rows_read": rows_read,
        "rows_imported": rows_imported,
        "entries_count": len(catalog["entries"]),
    }


def ensure_code_catalog(project_root):
    if os.path.exists(code_catalog_path(project_root)):
        raw_catalog = read_json(code_catalog_path(project_root), empty_code_catalog())
        normalized_catalog = _normalize_code_catalog(raw_catalog)
        if raw_catalog != normalized_catalog:
            save_code_catalog(project_root, normalized_catalog)
        return {"status": "success", "created": False, "catalog_path": code_catalog_path(project_root)}
    if os.path.exists(legacy_code_catalog_path(project_root)):
        save_code_catalog(project_root, read_json(legacy_code_catalog_path(project_root), empty_code_catalog()))
        return {"status": "success", "created": True, "catalog_path": code_catalog_path(project_root), "migrated": True}
    if os.path.isdir(code_source_dir(project_root)):
        result = import_codes_from_excel(project_root)
        result["created"] = result.get("status") == "success"
        return result
    write_json(code_catalog_path(project_root), empty_code_catalog())
    return {"status": "success", "created": True, "catalog_path": code_catalog_path(project_root)}


def code_catalog_map(project_root):
    ensure_code_catalog(project_root)
    code_map = {}
    for entry in load_code_catalog(project_root).get("entries", []):
        key = identity_key_from_values(entry.get("misura"), entry.get("spessore"), entry.get("materiale"), entry.get("finitura"))
        code = _entry_code(entry)
        code_map[key] = {
            **entry,
            "key": key,
            "codes": [code] if code else [],
            "code": code,
            "hasCode": bool(code),
        }
    return code_map


def get_codes_for_identity(project_root, misura, spessore, materiale, finitura):
    key = identity_key_from_values(misura, spessore, materiale, finitura)
    return code_catalog_map(project_root).get(key, {
        "key": key,
        "codes": [],
        "code": "",
        "hasCode": False,
    })


def set_code_catalog_entry(project_root, payload):
    payload = payload or {}
    key = identity_key_from_values(payload.get("misura"), payload.get("spessore"), payload.get("materiale"), payload.get("finitura"))
    identity = identity_dict_from_key(key)
    code_text = str(payload.get("code") if payload.get("code") is not None else payload.get("codes") or "").strip().upper()
    code_parts = [part.strip() for part in re.split(r"[,;/\n]+", code_text) if part.strip()]
    if len(code_parts) > 1:
        return {"status": "error", "message": "Ogni misura puo avere un solo codice."}
    code = code_parts[0] if code_parts else ""

    catalog = load_code_catalog(project_root)
    entries = []
    found = False
    for entry in catalog.get("entries", []):
        entry_key = identity_key_from_values(entry.get("misura"), entry.get("spessore"), entry.get("materiale"), entry.get("finitura"))
        if entry_key != key:
            entries.append(entry)
            continue
        found = True
        if code:
            entries.append({
                **entry,
                "key": key,
                "misura": identity["misura"],
                "spessore": identity["spessore"],
                "materiale": identity["materiale"],
                "finitura": identity["finitura"],
                "forma": identity["forma"],
                "code": code,
                "manual": True,
            })

    if not found and code:
        entries.append({
            "key": key,
            "misura": identity["misura"],
            "spessore": identity["spessore"],
            "materiale": identity["materiale"],
            "finitura": identity["finitura"],
            "forma": identity["forma"],
            "code": code,
            "sources": [],
            "manual": True,
        })

    catalog["entries"] = sorted(entries, key=lambda item: (
            item.get("forma", ""),
            item.get("misura", ""),
            _sort_number(item.get("spessore")),
            item.get("materiale", ""),
            item.get("finitura", ""),
        ))
    save_code_catalog(project_root, catalog)
    return {"status": "success", "message": "Codice aggiornato.", "code": code, "codes": [code] if code else []}


def backup_existing_database(project_root):
    db_dir = database_dir(project_root)
    if not os.path.isdir(db_dir):
        return None

    backup_root = os.path.join(os.path.dirname(db_dir), "_Backups")
    os.makedirs(backup_root, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    backup_dir = os.path.join(backup_root, f"Database Tubi backup {stamp}")
    shutil.copytree(db_dir, backup_dir)
    return backup_dir


def import_from_excel(project_root, excel_path, replace=True):
    if not excel_path or not os.path.exists(excel_path):
        return {
            "status": "error",
            "message": f"File Excel non trovato: {excel_path or ''}",
        }

    if replace:
        backup_existing_database(project_root)
        if os.path.isdir(database_dir(project_root)):
            shutil.rmtree(database_dir(project_root))

    os.makedirs(database_dir(project_root), exist_ok=True)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    locations = []
    total_rows = 0
    total_quantity = 0

    for sheet_name in wb.sheetnames:
        if not is_location_sheet(sheet_name):
            continue

        ws = wb[sheet_name]
        ripiani = parse_location_sheet(ws)
        location = {
            "schema_version": SCHEMA_VERSION,
            "name": sheet_name,
            "type": location_type(sheet_name),
            "ripiani": ripiani,
        }

        filename = safe_location_filename(sheet_name)
        write_json(os.path.join(database_dir(project_root), filename), location)
        locations.append({
            "name": sheet_name,
            "type": location["type"],
            "file": filename,
        })

        for ripiano in ripiani:
            for tube in ripiano.get("tubi", []):
                total_rows += 1
                total_quantity += int(tube.get("quantita") or 0)

    manifest = {
        "schema_version": SCHEMA_VERSION,
        "created_from": os.path.abspath(excel_path),
        "imported_at": datetime.now().isoformat(timespec="seconds"),
        "locations": locations,
    }
    write_json(manifest_path(project_root), manifest)

    if not os.path.exists(movements_path(project_root)):
        with open(movements_path(project_root), "w", encoding="utf-8"):
            pass

    return {
        "status": "success",
        "database_path": database_dir(project_root),
        "locations_count": len(locations),
        "rows_count": total_rows,
        "total_quantity": total_quantity,
    }


def database_exists(project_root):
    manifest = manifest_path(project_root)
    return os.path.exists(manifest) and os.path.isdir(database_dir(project_root))


def ensure_database(project_root, excel_path=None):
    if database_exists(project_root):
        return {"status": "success", "database_path": database_dir(project_root), "created": False}
    if excel_path and os.path.exists(excel_path):
        result = import_from_excel(project_root, excel_path, replace=True)
        result["created"] = result.get("status") == "success"
        return result
    return {
        "status": "missing",
        "message": "Database Tubi non trovato. Importa un file Excel per generarlo.",
        "database_path": database_dir(project_root),
    }


def load_manifest(project_root):
    return read_json(manifest_path(project_root), {"schema_version": SCHEMA_VERSION, "locations": []})


def load_location_file(project_root, file_name):
    path = os.path.join(database_dir(project_root), file_name)
    return read_json(path, None)


def save_location_file(project_root, file_name, data):
    write_json(os.path.join(database_dir(project_root), file_name), data)


def find_location_ref(project_root, location_name):
    for location_ref in load_manifest(project_root).get("locations", []):
        if str(location_ref.get("name", "")).lower() == str(location_name or "").lower():
            return location_ref
    return None


def ensure_ripiano(location, ripiano_name):
    ripiano_name = str(ripiano_name or "").strip() or "Ripiano 1"
    for ripiano in location.setdefault("ripiani", []):
        if str(ripiano.get("name", "")).lower() == ripiano_name.lower():
            ripiano.setdefault("tubi", [])
            return ripiano

    ripiano = {"name": ripiano_name, "tubi": []}
    location.setdefault("ripiani", []).append(ripiano)
    return ripiano


def load_location_by_name(project_root, location_name):
    location_ref = find_location_ref(project_root, location_name)
    if not location_ref:
        return None, None
    return location_ref, load_location_file(project_root, location_ref.get("file", ""))


def iter_position_tubes(project_root):
    manifest = load_manifest(project_root)
    for location_ref in manifest.get("locations", []):
        location = load_location_file(project_root, location_ref.get("file", ""))
        if not location:
            continue
        location_name = location.get("name") or location_ref.get("name") or ""
        location_type_name = location.get("type") or location_ref.get("type") or ""
        for ripiano in location.get("ripiani", []):
            ripiano_name = ripiano.get("name") or ""
            for tube in ripiano.get("tubi", []):
                yield location_name, location_type_name, ripiano_name, tube


def _spessore_candidates(spessore):
    sp = normalize_decimal_text(spessore)
    candidates = {sp}
    try:
        value = float(sp)
        if abs(value - 1.5) < 0.01:
            candidates.add("1.6")
        elif abs(value - 1.6) < 0.01:
            candidates.add("1.5")
    except ValueError:
        pass
    return candidates


def summarize_inventory(project_root, filters=None):
    filters = filters or {}
    grouped = {}
    codes_by_key = code_catalog_map(project_root)

    for location_name, location_type_name, ripiano_name, tube in iter_position_tubes(project_root):
        key = tube_key(tube)
        if not passes_filters(tube, filters):
            continue

        if key not in grouped:
            grouped[key] = {
                "misura": key[0],
                "spessore": decimal_number(key[1]),
                "materiale": key[2],
                "finitura": key[3],
                "forma": tube.get("forma") or classify_tube(key[0])[0],
                "dimensioni": tube.get("dimensioni") or classify_tube(key[0])[1],
                "quantitaTotale": 0,
                "posizioni": [],
            }

        qty = int(tube.get("quantita") or 0)
        grouped[key]["quantitaTotale"] += qty
        grouped[key]["posizioni"].append({
            "scaffalatura": location_name,
            "tipo": location_type_name,
            "ripiano": ripiano_name,
            "quantita": qty,
            "label": format_location_label(location_name, ripiano_name),
        })

    rows = []
    code_status = str(filters.get("codice") or "tutti").lower()
    for row in grouped.values():
        key = identity_key_from_values(row.get("misura"), row.get("spessore"), row.get("materiale"), row.get("finitura"))
        code_info = codes_by_key.get(key, {"codes": [], "code": "", "hasCode": False})
        row["codici"] = code_info.get("codes", [])
        row["codice"] = code_info.get("code", "")
        row["hasCode"] = bool(code_info.get("hasCode"))
        if code_status == "con" and not row["hasCode"]:
            continue
        if code_status == "senza" and row["hasCode"]:
            continue
        rows.append(row)

    rows.sort(key=lambda row: (
        row.get("forma", ""),
        str(row.get("misura", "")),
        float(row.get("spessore") or 0),
        str(row.get("materiale", "")),
        str(row.get("finitura", "")),
    ))
    return rows


def load_locations(project_root, filters=None):
    filters = filters or {}
    locations = []
    manifest = load_manifest(project_root)
    codes_by_key = code_catalog_map(project_root)
    code_status = str(filters.get("codice") or "tutti").lower()

    for location_ref in manifest.get("locations", []):
        location = load_location_file(project_root, location_ref.get("file", ""))
        if not location:
            continue

        filtered_location = {
            "name": location.get("name") or location_ref.get("name"),
            "type": location.get("type") or location_ref.get("type"),
            "file": location_ref.get("file"),
            "ripiani": [],
        }

        for ripiano in location.get("ripiani", []):
            tubes = []
            for tube in ripiano.get("tubi", []):
                if not passes_filters(tube, filters):
                    continue
                key = identity_key_from_tube(tube)
                code_info = codes_by_key.get(key, {"codes": [], "code": "", "hasCode": False})
                if code_status == "con" and not code_info.get("hasCode"):
                    continue
                if code_status == "senza" and code_info.get("hasCode"):
                    continue
                enriched = {
                    **tube,
                    "codici": code_info.get("codes", []),
                    "codice": code_info.get("code", ""),
                    "hasCode": bool(code_info.get("hasCode")),
                }
                tubes.append(enriched)
            filtered_location["ripiani"].append({
                "name": ripiano.get("name") or "",
                "tubi": sort_tubes(tubes),
            })

        locations.append(filtered_location)

    return locations


def get_location_targets(project_root):
    targets = []
    manifest = load_manifest(project_root)

    for location_ref in manifest.get("locations", []):
        location = load_location_file(project_root, location_ref.get("file", ""))
        if not location:
            continue
        targets.append({
            "name": location.get("name") or location_ref.get("name"),
            "type": location.get("type") or location_ref.get("type"),
            "ripiani": [ripiano.get("name") or "" for ripiano in location.get("ripiani", [])],
        })

    return targets


def location_is_empty(location):
    for ripiano in location.get("ripiani", []):
        for tube in ripiano.get("tubi", []):
            if int(tube.get("quantita") or 0) > 0:
                return False
    return True


def ripiano_is_empty(ripiano):
    return all(int(tube.get("quantita") or 0) <= 0 for tube in ripiano.get("tubi", []))


def save_manifest(project_root, manifest):
    write_json(manifest_path(project_root), manifest)


def add_scaffolding(project_root, name):
    name = str(name or "").strip()
    if not name:
        return {"status": "error", "message": "Nome scaffalatura mancante."}

    manifest = load_manifest(project_root)
    if any(str(item.get("name", "")).lower() == name.lower() for item in manifest.get("locations", [])):
        return {"status": "error", "message": f"Scaffalatura già esistente: {name}"}

    file_name = safe_location_filename(name)
    existing_files = {str(item.get("file", "")).lower() for item in manifest.get("locations", [])}
    if file_name.lower() in existing_files or os.path.exists(os.path.join(database_dir(project_root), file_name)):
        return {"status": "error", "message": f"File già esistente: {file_name}"}

    location = {
        "schema_version": SCHEMA_VERSION,
        "name": name,
        "type": "scaffalatura",
        "ripiani": [{"name": "Ripiano 1", "tubi": []}],
    }
    write_json(os.path.join(database_dir(project_root), file_name), location)
    manifest.setdefault("locations", []).append({"name": name, "type": "scaffalatura", "file": file_name})
    save_manifest(project_root, manifest)
    append_movement(project_root, {"action": "add_scaffolding", "location": name})
    return {"status": "success", "message": "Scaffalatura aggiunta."}


def remove_scaffolding(project_root, name):
    name = str(name or "").strip()
    manifest = load_manifest(project_root)
    location_ref = next((item for item in manifest.get("locations", []) if str(item.get("name", "")).lower() == name.lower()), None)
    if not location_ref:
        return {"status": "error", "message": f"Scaffalatura non trovata: {name}"}
    if str(location_ref.get("type", "")).lower() == "terra":
        return {"status": "error", "message": "Terra non può essere rimossa."}

    location = load_location_file(project_root, location_ref.get("file", ""))
    if not location:
        return {"status": "error", "message": "File scaffalatura non leggibile."}
    if not location_is_empty(location):
        return {"status": "error", "message": "La scaffalatura contiene ancora tubi. Svuotala prima di rimuoverla."}

    manifest["locations"] = [item for item in manifest.get("locations", []) if item is not location_ref]
    save_manifest(project_root, manifest)
    try:
        os.remove(os.path.join(database_dir(project_root), location_ref.get("file", "")))
    except OSError:
        pass
    append_movement(project_root, {"action": "remove_scaffolding", "location": name})
    return {"status": "success", "message": "Scaffalatura rimossa."}


def add_ripiano(project_root, location_name, ripiano_name):
    location_ref, location = load_location_by_name(project_root, location_name)
    if not location_ref or not location:
        return {"status": "error", "message": f"Ubicazione non trovata: {location_name}"}

    ripiano_name = str(ripiano_name or "").strip()
    if not ripiano_name:
        existing_numbers = []
        for ripiano in location.get("ripiani", []):
            match = re.search(r"(?i)ripiano\s*(\d+)", str(ripiano.get("name", "")))
            if match:
                existing_numbers.append(int(match.group(1)))
        ripiano_name = f"Ripiano {(max(existing_numbers) + 1) if existing_numbers else 1}"

    if any(str(ripiano.get("name", "")).lower() == ripiano_name.lower() for ripiano in location.get("ripiani", [])):
        return {"status": "error", "message": f"Ripiano già esistente: {ripiano_name}"}

    location.setdefault("ripiani", []).append({"name": ripiano_name, "tubi": []})
    save_location_file(project_root, location_ref["file"], location)
    append_movement(project_root, {"action": "add_ripiano", "target": {"location": location_name, "ripiano": ripiano_name}})
    return {"status": "success", "message": "Ripiano aggiunto."}


def remove_ripiano(project_root, location_name, ripiano_name):
    location_ref, location = load_location_by_name(project_root, location_name)
    if not location_ref or not location:
        return {"status": "error", "message": f"Ubicazione non trovata: {location_name}"}

    ripiani = location.get("ripiani", [])
    ripiano = next((item for item in ripiani if str(item.get("name", "")).lower() == str(ripiano_name or "").lower()), None)
    if not ripiano:
        return {"status": "error", "message": f"Ripiano non trovato: {ripiano_name}"}
    if not ripiano_is_empty(ripiano):
        return {"status": "error", "message": "Il ripiano contiene ancora tubi. Svuotalo prima di rimuoverlo."}
    if len(ripiani) <= 1:
        return {"status": "error", "message": "Una ubicazione deve avere almeno un ripiano/zona."}

    location["ripiani"] = [item for item in ripiani if item is not ripiano]
    save_location_file(project_root, location_ref["file"], location)
    append_movement(project_root, {"action": "remove_ripiano", "source": {"location": location_name, "ripiano": ripiano_name}})
    return {"status": "success", "message": "Ripiano rimosso."}


def passes_filters(tube, filters):
    forma = str(filters.get("forma") or "tutti").lower()
    materiale = str(filters.get("materiale") or "tutti").upper()
    finitura = str(filters.get("finitura") or "tutti").upper()
    misura = normalize_misura(filters.get("misura"))
    spessore = normalize_decimal_text(filters.get("spessore"))

    tube_forma = str(tube.get("forma") or "").lower()
    if forma != "tutti":
        if forma == "tubolari" and tube_forma not in {"quadrato", "rettangolare"}:
            return False
        if forma in {"tondo", "quadrato", "rettangolare"} and tube_forma != forma:
            return False

    if materiale != "TUTTI" and normalize_materiale(tube.get("materiale")) != materiale:
        return False
    if finitura != "TUTTI" and normalize_finitura(tube.get("finitura")) != finitura:
        return False
    if misura and misura not in normalize_misura(tube.get("misura")):
        return False
    if spessore and normalize_decimal_text(tube.get("spessore")) not in _spessore_candidates(spessore):
        return False

    return True


def format_location_label(location_name, ripiano_name):
    location_name = str(location_name or "").strip()
    ripiano_name = str(ripiano_name or "").strip()
    if not ripiano_name:
        return location_name

    match = re.search(r"(?i)scaffalatura\s+(.+)", location_name)
    scaff_code = f"S{match.group(1).strip().replace(' ', '')}" if match else location_name

    ripiano_match = re.search(r"(?i)ripiano\s*(\d+)", ripiano_name)
    ripiano_code = f"R{ripiano_match.group(1)}" if ripiano_match else ripiano_name

    if location_name.lower() == "terra":
        return ripiano_name
    return f"{scaff_code} {ripiano_code}".strip()


def get_locations_for_tube(project_root, misura, spessore, materiale, finitura):
    misura_norm = normalize_misura(misura)
    material_norm = normalize_materiale(materiale)
    finish_norm = normalize_finitura(finitura)
    spessore_set = _spessore_candidates(spessore)

    merged = defaultdict(int)
    for location_name, _, ripiano_name, tube in iter_position_tubes(project_root):
        key = tube_key(tube)
        if key[0] == misura_norm and key[2] == material_norm and key[3] == finish_norm and key[1] in spessore_set:
            qty = int(tube.get("quantita") or 0)
            if qty > 0:
                merged[format_location_label(location_name, ripiano_name)] += qty

    return [(label, merged[label]) for label in sorted(merged.keys())]


def get_position_records_for_tube(project_root, misura, spessore, materiale, finitura):
    misura_norm = normalize_misura(misura)
    material_norm = normalize_materiale(materiale)
    finish_norm = normalize_finitura(finitura)
    spessore_set = _spessore_candidates(spessore)

    positions = []
    for location_name, location_type_name, ripiano_name, tube in iter_position_tubes(project_root):
        key = tube_key(tube)
        qty = int(tube.get("quantita") or 0)
        if qty <= 0:
            continue
        if key[0] == misura_norm and key[2] == material_norm and key[3] == finish_norm and key[1] in spessore_set:
            positions.append({
                "label": format_location_label(location_name, ripiano_name),
                "source_location": location_name,
                "source_ripiano": ripiano_name,
                "location_type": location_type_name,
                "misura": key[0],
                "spessore": decimal_number(key[1]),
                "materiale": key[2],
                "finitura": key[3],
                "quantita": qty,
            })

    positions.sort(key=lambda item: str(item.get("label", "")).lower())
    return positions


def database_debug(project_root):
    manifest = load_manifest(project_root)
    location_count = len(manifest.get("locations", []))
    position_rows = 0
    total_quantity = 0
    for _, _, _, tube in iter_position_tubes(project_root):
        position_rows += 1
        total_quantity += int(tube.get("quantita") or 0)

    return {
        "database_path": database_dir(project_root),
        "manifest_path": manifest_path(project_root),
        "exists": database_exists(project_root),
        "schema_version": manifest.get("schema_version"),
        "imported_at": manifest.get("imported_at"),
        "created_from": manifest.get("created_from"),
        "locations_count": location_count,
        "position_rows": position_rows,
        "total_quantity": total_quantity,
    }


def append_movement(project_root, event):
    os.makedirs(database_dir(project_root), exist_ok=True)
    event = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        **event,
    }
    with open(movements_path(project_root), "a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")


def tube_from_payload(payload):
    return make_tube(
        payload.get("misura"),
        payload.get("spessore"),
        payload.get("materiale"),
        payload.get("finitura"),
        payload.get("quantita", 0),
    )


def parse_positive_quantity(payload):
    quantity = parse_quantity(payload.get("quantita"))
    if quantity <= 0:
        raise ValueError("La quantità deve essere maggiore di zero.")
    return quantity


def add_tube(project_root, payload):
    quantity = parse_positive_quantity(payload)
    location_name = payload.get("target_location") or payload.get("location")
    ripiano_name = payload.get("target_ripiano") or payload.get("ripiano")
    location_ref, location = load_location_by_name(project_root, location_name)
    if not location_ref or not location:
        return {"status": "error", "message": f"Ubicazione non trovata: {location_name}"}

    tube = tube_from_payload({**payload, "quantita": quantity})
    ripiano = ensure_ripiano(location, ripiano_name)
    merge_tube_into_list(ripiano["tubi"], tube)
    ripiano["tubi"] = sort_tubes(ripiano["tubi"])
    save_location_file(project_root, location_ref["file"], location)
    append_movement(project_root, {
        "action": "add",
        "target": {"location": location_name, "ripiano": ripiano.get("name")},
        "tube": tube,
    })
    return {"status": "success", "message": "Tubo aggiunto."}


def remove_tube_from_ripiano(ripiano, tube_identity, quantity):
    identity_key = tube_key(tube_identity)
    for index, existing in enumerate(ripiano.get("tubi", [])):
        if tube_key(existing) != identity_key:
            continue

        available = int(existing.get("quantita") or 0)
        if quantity > available:
            raise ValueError(f"Quantità richiesta ({quantity}) maggiore di quella disponibile ({available}).")

        existing["quantita"] = available - quantity
        if existing["quantita"] <= 0:
            ripiano["tubi"].pop(index)
        return {
            **tube_identity,
            "quantita": quantity,
        }

    raise ValueError("Tubo non trovato nella posizione selezionata.")


def remove_tube(project_root, payload):
    quantity = parse_positive_quantity(payload)
    location_name = payload.get("source_location") or payload.get("location")
    ripiano_name = payload.get("source_ripiano") or payload.get("ripiano")
    location_ref, location = load_location_by_name(project_root, location_name)
    if not location_ref or not location:
        return {"status": "error", "message": f"Ubicazione non trovata: {location_name}"}

    tube_identity = tube_from_payload({**payload, "quantita": quantity})

    try:
        ripiano = ensure_ripiano(location, ripiano_name)
        removed = remove_tube_from_ripiano(ripiano, tube_identity, quantity)
        ripiano["tubi"] = sort_tubes(ripiano["tubi"])
        save_location_file(project_root, location_ref["file"], location)
        append_movement(project_root, {
            "action": "remove",
            "source": {"location": location_name, "ripiano": ripiano.get("name")},
            "tube": removed,
        })
        return {"status": "success", "message": "Tubo rimosso."}
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}


def set_tube_quantity(project_root, payload):
    quantity = parse_quantity(payload.get("quantita"))
    if quantity < 0:
        return {"status": "error", "message": "La quantita non puo essere negativa."}

    location_name = payload.get("source_location") or payload.get("location")
    ripiano_name = payload.get("source_ripiano") or payload.get("ripiano")
    location_ref, location = load_location_by_name(project_root, location_name)
    if not location_ref or not location:
        return {"status": "error", "message": f"Ubicazione non trovata: {location_name}"}

    tube_identity = tube_from_payload({**payload, "quantita": quantity})
    identity_key = tube_key(tube_identity)
    ripiano = ensure_ripiano(location, ripiano_name)

    for index, existing in enumerate(ripiano.get("tubi", [])):
        if tube_key(existing) != identity_key:
            continue

        old_quantity = int(existing.get("quantita") or 0)
        if quantity <= 0:
            ripiano["tubi"].pop(index)
        else:
            existing["quantita"] = quantity

        ripiano["tubi"] = sort_tubes(ripiano["tubi"])
        save_location_file(project_root, location_ref["file"], location)
        append_movement(project_root, {
            "action": "set_quantity",
            "source": {"location": location_name, "ripiano": ripiano.get("name")},
            "tube": {**tube_identity, "quantita": quantity},
            "old_quantity": old_quantity,
            "new_quantity": quantity,
            "delta": quantity - old_quantity,
        })
        return {"status": "success", "message": "Quantita aggiornata."}

    return {"status": "error", "message": "Tubo non trovato nella posizione selezionata."}


def move_tube(project_root, payload):
    quantity = parse_positive_quantity(payload)
    source_location_name = payload.get("source_location")
    source_ripiano_name = payload.get("source_ripiano")
    target_location_name = payload.get("target_location")
    target_ripiano_name = payload.get("target_ripiano")

    source_ref, source_location = load_location_by_name(project_root, source_location_name)
    target_ref, target_location = load_location_by_name(project_root, target_location_name)
    if not source_ref or not source_location:
        return {"status": "error", "message": f"Ubicazione origine non trovata: {source_location_name}"}
    if not target_ref or not target_location:
        return {"status": "error", "message": f"Ubicazione destinazione non trovata: {target_location_name}"}

    if source_ref["file"] == target_ref["file"]:
        target_location = source_location

    tube_identity = tube_from_payload({**payload, "quantita": quantity})

    try:
        source_ripiano = ensure_ripiano(source_location, source_ripiano_name)
        removed = remove_tube_from_ripiano(source_ripiano, tube_identity, quantity)
        source_ripiano["tubi"] = sort_tubes(source_ripiano["tubi"])

        target_ripiano = ensure_ripiano(target_location, target_ripiano_name)
        merge_tube_into_list(target_ripiano["tubi"], removed)
        target_ripiano["tubi"] = sort_tubes(target_ripiano["tubi"])

        if source_ref["file"] == target_ref["file"]:
            save_location_file(project_root, source_ref["file"], source_location)
        else:
            save_location_file(project_root, source_ref["file"], source_location)
            save_location_file(project_root, target_ref["file"], target_location)

        append_movement(project_root, {
            "action": "move",
            "source": {"location": source_location_name, "ripiano": source_ripiano.get("name")},
            "target": {"location": target_location_name, "ripiano": target_ripiano.get("name")},
            "tube": removed,
        })
        return {"status": "success", "message": "Tubo spostato."}
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}


def _inventory_totals(project_root):
    totals = defaultdict(int)
    for _, _, _, tube in iter_position_tubes(project_root):
        key = identity_key_from_tube(tube)
        totals[key] += int(tube.get("quantita") or 0)
    return dict(totals)


def load_adhoc_state(project_root):
    state = read_json(adhoc_state_path(project_root), {})
    state.setdefault("schema_version", SCHEMA_VERSION)
    state.setdefault("baseline_at", None)
    state.setdefault("baseline", {})
    return state


def _adhoc_rows_from_delta(project_root, baseline, current):
    codes_by_key = code_catalog_map(project_root)
    rows = []

    def delta_label(delta_value):
        amount = abs(int(delta_value or 0))
        if delta_value < 0:
            return "1 verga usata" if amount == 1 else f"{amount} verghe usate"
        return "1 verga aggiunta" if amount == 1 else f"{amount} verghe aggiunte"

    for key in sorted(set(baseline) | set(current)):
        old_qty = int(baseline.get(key, 0) or 0)
        new_qty = int(current.get(key, 0) or 0)
        delta = new_qty - old_qty
        if delta == 0:
            continue
        identity = identity_dict_from_key(key)
        code_info = codes_by_key.get(key, {"codes": [], "code": "", "hasCode": False})
        rows.append({
            **identity,
            "misuraTubo": display_tube_measure(identity["misura"], identity["spessore"], identity["materiale"], identity["finitura"], include_defaults=True),
            "codice": code_info.get("code") or "CODICE MANCANTE",
            "codici": code_info.get("codes", []),
            "hasCode": bool(code_info.get("hasCode")),
            "quantita": new_qty,
            "quantitaLabel": str(new_qty),
            "delta": delta,
            "deltaLabel": delta_label(delta),
        })
    return rows


def get_adhoc_report_data(project_root):
    ensure_code_catalog(project_root)
    state = load_adhoc_state(project_root)
    if not state.get("baseline_at"):
        return {
            "status": "missing_baseline",
            "message": "Baseline ADHOC non impostata. Premi Azzera ADHOC dopo aver aggiornato il database reale.",
            "rows": [],
        }

    current = _inventory_totals(project_root)
    rows = _adhoc_rows_from_delta(project_root, state.get("baseline", {}), current)
    return {
        "status": "success",
        "baseline_at": state.get("baseline_at"),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "rows": rows,
    }


def reset_adhoc_baseline(project_root):
    ensure_code_catalog(project_root)
    old_state = load_adhoc_state(project_root)
    current = _inventory_totals(project_root)
    reset_at = datetime.now().isoformat(timespec="seconds")
    previous_rows = []
    if old_state.get("baseline_at"):
        previous_rows = _adhoc_rows_from_delta(project_root, old_state.get("baseline", {}), current)

    os.makedirs(database_dir(project_root), exist_ok=True)
    with open(adhoc_log_path(project_root), "a", encoding="utf-8") as f:
        f.write(json.dumps({
            "timestamp": reset_at,
            "previous_baseline_at": old_state.get("baseline_at"),
            "new_baseline_at": reset_at,
            "rows": previous_rows,
        }, ensure_ascii=False) + "\n")

    state = {
        "schema_version": SCHEMA_VERSION,
        "baseline_at": reset_at,
        "baseline": current,
    }
    write_json(adhoc_state_path(project_root), state)
    return {
        "status": "success",
        "message": "Baseline ADHOC aggiornata.",
        "baseline_at": reset_at,
        "rows_logged": len(previous_rows),
        "state_path": adhoc_state_path(project_root),
        "log_path": adhoc_log_path(project_root),
    }


def read_summary_sheet(excel_path, sheet_name):
    if not excel_path or not os.path.exists(excel_path):
        return {}

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return {}

    ws = wb[sheet_name]
    summary = {}
    for row in range(2, (ws.max_row or 0) + 1):
        misura = normalize_misura(ws.cell(row=row, column=1).value)
        spessore = normalize_decimal_text(ws.cell(row=row, column=2).value)
        materiale = normalize_materiale(ws.cell(row=row, column=3).value)
        finitura = normalize_finitura(ws.cell(row=row, column=4).value)
        quantita = parse_quantity(ws.cell(row=row, column=5).value)
        if misura and materiale and finitura:
            summary[(misura, spessore, materiale, finitura)] = quantita
    return summary


def validate_against_excel_summaries(project_root, excel_path):
    if not excel_path or not os.path.exists(excel_path):
        return {"status": "error", "message": "File Excel non trovato."}

    db_summary = {}
    for row in summarize_inventory(project_root):
        key = (
            normalize_misura(row["misura"]),
            normalize_decimal_text(row["spessore"]),
            normalize_materiale(row["materiale"]),
            normalize_finitura(row["finitura"]),
        )
        db_summary[key] = int(row.get("quantitaTotale") or 0)

    excel_summary = {}
    excel_summary.update(read_summary_sheet(excel_path, "ConteggioTubiTondi"))
    excel_summary.update(read_summary_sheet(excel_path, "ConteggioTubolari"))

    mismatches = []
    for key in sorted(set(db_summary) | set(excel_summary)):
        db_qty = db_summary.get(key, 0)
        excel_qty = excel_summary.get(key, 0)
        if db_qty != excel_qty:
            mismatches.append({
                "misura": key[0],
                "spessore": key[1],
                "materiale": key[2],
                "finitura": key[3],
                "database": db_qty,
                "excel": excel_qty,
            })

    return {
        "status": "success",
        "mismatches": mismatches,
        "database_rows": len(db_summary),
        "excel_rows": len(excel_summary),
    }
